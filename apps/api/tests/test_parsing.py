import io

import pytest
from docx import Document

from app.rag.parsing import extract_text
from app.rag.parsing.docx import extract_text_from_docx


def _build_docx_bytes(paragraphs: list[str]) -> bytes:
    document = Document()
    for p in paragraphs:
        document.add_paragraph(p)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_extract_text_from_docx_round_trip():
    file_bytes = _build_docx_bytes(["Jane Doe", "Software Engineer", "jane@example.com"])
    text = extract_text_from_docx(file_bytes)
    assert "Jane Doe" in text
    assert "Software Engineer" in text
    assert "jane@example.com" in text


def test_extract_text_from_docx_skips_empty_paragraphs():
    file_bytes = _build_docx_bytes(["Line one", "", "   ", "Line two"])
    text = extract_text_from_docx(file_bytes)
    assert text == "Line one\nLine two"


def test_extract_text_dispatches_to_docx_by_extension():
    file_bytes = _build_docx_bytes(["Dispatch check"])
    text = extract_text("resume.DOCX", file_bytes)
    assert "Dispatch check" in text


def test_extract_text_raises_for_unsupported_extension():
    with pytest.raises(ValueError, match="Unsupported file type"):
        extract_text("resume.txt", b"plain text")


def test_pdf_text_keeps_words_apart_and_drops_icon_placeholders(monkeypatch):
    # LaTeX resumes set spaces tightly; the default tolerance glued words together
    # ("BachelorofTechnologyinComputerScience") and broke keyword matching.
    from app.rag.parsing import pdf as pdf_module

    seen = {}

    class FakePage:
        def extract_text(self, **kwargs):
            seen.update(kwargs)
            return "(cid:131) +91 6301125158\nBachelor of Technology in Computer Science"

    class FakePdf:
        pages = [FakePage()]

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

    monkeypatch.setattr(pdf_module.pdfplumber, "open", lambda _f: FakePdf())
    text = pdf_module.extract_text_from_pdf(b"%PDF")
    assert seen == {"x_tolerance": 1.5}
    assert "(cid:" not in text and "Computer Science" in text
